from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "submission-archive" / "mvp-report" / "Semantic_Query_Engine_MVP_and_Production_Report.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
GRAY = RGBColor(89, 89, 89)


def set_font(run, size=None, color=None, bold=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    set_font(run, 9.5, color, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for col in list(grid.gridCol_lst):
        grid.remove(col)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(grid_col)

    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
            tc_pr = row.cells[index]._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    run = p.add_run(text)
    set_font(run, 16 if level == 1 else 12.5, BLUE if level == 1 else DARK_BLUE, True)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        set_font(run, 11, None, True)
        remainder = text[len(bold_lead):]
        run = p.add_run(remainder)
        set_font(run, 11)
    else:
        run = p.add_run(text)
        set_font(run, 11)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(item)
        set_font(run, 10.5)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, True)
        shade(table.rows[0].cells[i], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("Semantic Query Engine | MVP and Production Readiness Report")
    set_font(header_run, 9, GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Semantic Query Engine")
    set_font(footer_run, 9, GRAY)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("Semantic Query Engine")
    set_font(title_run, 24, DARK_BLUE, True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run("MVP Delivery, Production-Grade Target State, and Delivery Roadmap")
    set_font(subtitle_run, 13, GRAY)

    add_table(
        doc,
        ["Report purpose", "Scope"],
        [["Document the current MVP and define the work required for a governed production release.", "Natural-language FMCG analytics across sales, promotion, inventory, SKU, category, channel, and trend questions."]],
        [2.0, 4.5],
    )

    add_heading(doc, "1. Executive Summary")
    add_body(doc, "Semantic Query Engine is an AI-assisted FMCG analytics application that lets business users ask natural-language questions and receive governed SQL-backed answers. The MVP demonstrates a complete analytical path: question intake, intent planning, semantic retrieval, SQL generation, validation, execution in DuckDB, structured explanation, visualization, and downloadable data.")
    add_body(doc, "The MVP is appropriate for demonstrations, assessment, and controlled internal prototyping. A production release should retain the same core flow but replace local data access and heuristic safeguards with governed warehouse access, identity-aware authorization, stronger evaluation, observability, and operational controls.")

    add_heading(doc, "2. Current MVP: What Has Been Developed")
    add_table(
        doc,
        ["Capability", "MVP implementation", "Business value"],
        [
            ["Conversational analytics", "Streamlit chat interface with an analysis catalog, recent-query shortcuts, period selection, and CSV export.", "Business users can explore data without writing SQL."],
            ["Question planning", "Planner classifies descriptive, comparative, diagnostic, and ambiguous requests and requests clarification when needed.", "Reduces accidental interpretation of vague business questions."],
            ["Semantic grounding", "Semantic metadata describes tables, columns, and certified business metrics; retriever selects relevant context.", "Connects language to business meaning rather than raw schema alone."],
            ["Text-to-SQL", "LLM-backed generation with deterministic fallback templates for revenue, promotion, stock, SKU, category, channel, brand, and trends.", "Keeps core demo flows usable even when an LLM is unavailable."],
            ["Safety and execution", "SQLGlot parser validation, allowed-table controls, schema checks, metric checks, repair attempts, and DuckDB execution.", "Improves safety, traceability, and correctness before results are shown."],
            ["Business response", "Narrative summary, key metric, comparison context, interactive chart controls, data table, and SQL/agent trace.", "Makes analytical output understandable and auditable."],
            ["Quality controls", "Gold evaluation cases and automated tests cover core flows, validator behavior, period filtering, and analysis-catalog queries.", "Provides regression protection during iteration."],
        ],
        [1.35, 3.2, 1.95],
    )

    add_heading(doc, "3. MVP Architecture")
    add_body(doc, "The application is organized as a stateful, multi-stage pipeline. Each stage has a specific responsibility, which makes the flow easier to test and explain than a single unrestricted LLM call.")
    add_bullets(doc, [
        "User interface: Streamlit accepts questions, offers analysis templates, applies an optional year filter, caches repeated work, and presents charts and tables.",
        "Planner: extracts entities and selects an intent archetype; ambiguous questions trigger a clarification response.",
        "Schema retriever: loads semantic metadata and supplies relevant table, column, and business-metric context.",
        "SQL generator: uses an LLM when configured and controlled fallback templates otherwise.",
        "Validator and repair loop: parses SQL, checks its structure and references, validates approved metrics, and requests bounded repair when validation fails.",
        "Execution and synthesis: executes approved SQL against DuckDB and converts the result into a concise business response.",
    ])

    add_heading(doc, "4. MVP Scope and Known Limitations")
    add_table(
        doc,
        ["Area", "Current MVP position", "Limitation"],
        [
            ["Data platform", "Local CSV files loaded into DuckDB.", "Appropriate for a prototype; not a managed enterprise data environment."],
            ["Access control", "No user identity or authorization model.", "No row-level security, column masking, or audit trail by user."],
            ["Semantic retrieval", "Metadata-driven retrieval with business definitions and keyword/synonym matching.", "Not yet an embedding-backed vector retrieval service with formal ownership and change control."],
            ["SQL coverage", "Strong templates for common FMCG questions plus LLM generation.", "Fallback paths remain heuristic and dataset-specific for unusual phrasing or complex joins."],
            ["Validation", "Parser, schema, and selected metric checks; bounded repair loop.", "Cannot guarantee every SQL query is semantically correct without broader test coverage and approved query patterns."],
            ["Operations", "Local caching and automated tests.", "No production telemetry, service-level objectives, alerting, or cost controls."],
        ],
        [1.25, 2.85, 2.4],
    )

    add_heading(doc, "5. Production-Grade Target Architecture")
    add_body(doc, "A production release should preserve the user-facing experience while changing where trust is enforced. The data platform, not the LLM, must be the authority for access, governance, and aggregation.")
    add_table(
        doc,
        ["Layer", "Production-grade capability", "Outcome"],
        [
            ["Identity and governance", "OAuth2 or SSO session identity; warehouse RBAC; row-level security and dynamic column masking.", "Every answer is limited to the user's approved data scope."],
            ["Data platform", "Snowflake or Databricks/Unity Catalog; curated tables, dbt transformations, materialized business metrics.", "Reliable, scalable, and governed analytical execution."],
            ["Semantic layer", "Versioned metric catalog, documented joins, metric ownership, embeddings/vector retrieval, and approved synonyms.", "Consistent business definitions and better context retrieval."],
            ["Query safety", "AST policy enforcement, live catalog validation, cost/scan/time limits, query tagging, parameterized filters, and allowlisted joins.", "Prevents unsafe, expensive, or semantically unsupported queries."],
            ["LLM operations", "Model routing, prompt/version management, response caching, rate limits, and fallback policy.", "Predictable cost, latency, and quality."],
            ["Observability", "Structured logs, traces, audit records, user feedback, quality dashboards, and alerting.", "Ability to detect failures, misuse, drift, and declining answer quality."],
            ["Human controls", "Approval workflow for high-stakes metrics and escalation to analysts when confidence is low.", "Appropriate safeguards for executive or financial decisions."],
        ],
        [1.25, 3.15, 2.1],
    )

    add_heading(doc, "6. Production Delivery Roadmap")
    add_table(
        doc,
        ["Phase", "Priority deliverables", "Exit criteria"],
        [
            ["1. Stabilize MVP", "Resolve remaining UI text/encoding issues; consolidate package/import layout; expand gold evaluation set; add LLM-mocked integration tests.", "All supported catalog questions are tested and no known user-facing errors remain."],
            ["2. Govern data access", "Move to a governed warehouse; implement SSO, RBAC, RLS, masking, query tagging, and audit logging.", "Data access is identity-aware and policy-enforced by the platform."],
            ["3. Strengthen accuracy", "Version semantic metrics and joins; build an embedding-backed retriever; add semantic-result tests and confidence thresholds.", "Measured accuracy meets agreed business thresholds on a representative benchmark."],
            ["4. Operate at scale", "Introduce service APIs, asynchronous jobs where needed, caching, monitoring, alerting, cost controls, and incident runbooks.", "Defined reliability, latency, and cost targets are consistently met."],
            ["5. Expand adoption", "Role-specific dashboards, feedback loops, analyst escalation, and additional data domains.", "Usage and feedback demonstrate repeatable value across business teams."],
        ],
        [1.05, 3.5, 1.95],
    )

    add_heading(doc, "7. Recommended Success Measures")
    add_bullets(doc, [
        "Analytical quality: execution success rate, semantic accuracy on gold questions, and clarification success rate.",
        "Business value: reduction in analyst turnaround time, self-service query adoption, and repeat use by business users.",
        "Reliability: p50/p95 response latency, query failure rate, data freshness, and service availability.",
        "Governance: percentage of queries executed under identity-aware policy, audit coverage, and blocked unsafe-query rate.",
        "Cost: model cost per successful answer, warehouse cost per query, cache hit rate, and high-cost query frequency.",
    ])

    add_heading(doc, "8. Conclusion")
    add_body(doc, "Semantic Query Engine demonstrates the essential mechanics of a trustworthy conversational analytics assistant: grounding in business metadata, structured planning, safe SQL execution, repair, explainability, and analyst-friendly output. The production path is not a rewrite; it is a controlled evolution from local prototype components to governed enterprise services, with security, measurement, and semantic quality made explicit at every stage.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
